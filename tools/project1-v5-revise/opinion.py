# 生成《修改意见与处理说明》
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = sys.argv[1]
C = RGBColor(0x15, 0x3A, 0x6B)
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)
st = d.styles['Normal']
st.font.name = '宋体'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.3
for name, size in (('Title', 18), ('Heading 1', 14), ('Heading 2', 12)):
    s = d.styles[name]
    s.font.name = '黑体'
    s.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    s.font.size = Pt(size)
    s.font.color.rgb = C
    s.font.bold = True


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hexcolor)
    tcpr.append(sh)


def table(head, rows, widths):
    t = d.add_table(rows=1, cols=len(head))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        shade(c, 'E8EEF7')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5)
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t


def p(text, bold_lead=None):
    para = d.add_paragraph()
    if bold_lead:
        para.add_run(bold_lead).bold = True
    para.add_run(text)
    return para


t = d.add_paragraph(style='Title')
t.add_run('项目一《搭建个人专属档案》修改意见与处理说明')
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
m = d.add_paragraph()
m.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = m.add_run('底稿：冯丽娜批注稿（2026 年 10 月 11 日，7 条批注）　修订：2026 年 10 月 12 日　修订稿：v5')
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

d.add_heading('一、总体意见', level=1)
p('新稿把原来的三个任务合并为两个，“迎新系统档案卡片”的情境线索贯穿始终，结构比 v4 更紧凑，可以作为定稿底稿。因此本轮以作者新稿为底稿修改，没有回到 v4。')
p('本轮按 7 条批注逐条修改，另外修正了合并任务后留下的编号、引用和前后不一致等 20 处问题。作者的写法和版式都保留；新增内容沿用原有样式。修订稿中每条批注下都有“已处理”的回复，审阅时可逐条核对，确认后在 Word 中“删除所有批注”即可。清洁版是删去全部批注的同一份文件。')

d.add_heading('二、批注逐条处理', level=1)
table(['序号', '位置', '批注要点', '处理'], [
    ['1', '三维目标', '根据调整好的内容重新调整', '按两个任务重写：知识目标 6 条、能力目标 6 条，素养目标不变。原“掌握 printf 输出各类数据”超出本项目所学（只学了输出文字和 %d），改为“会用 printf 输出文字和整数”；能力目标补“核对 AI 的回答”和“四要素提问法”。'],
    ['2', '知识导图（图1-2）', '根据调整好的内容重新调整', '重画为“任务一、任务二、AI升华”三支，节点与正文小标题一一对应。'],
    ['3', '图1-3', '左侧改“操作说明：手机连校园网”，右侧改“程序：新生专属档案”', '按批注文字重画，沿用全书配色；正文“按照图1-3的操作说明……”与新图一致。'],
    ['4', '图1-5 应用领域', '保留版式，用实物照片；加一两个大学生感兴趣的例子（游戏、点餐 App）', '① 图1-5 改为 2×3 的“照片＋标题＋一句说明”表格，文字可直接改。六格都换成了真实图片：五张取自开源项目公开发布的实拍照片，授权允许出版使用；系统软件一格是编者在 Linux 上的终端实拍截图。图下加了“图片来源”说明，来源和许可证明细见第四部分。如学校有实训室实拍照片，右键“更改图片”即可替换。② 新增图1-6：点餐小程序和贪吃蛇小游戏的真实运行画面，各配三步“怎么做出来的”，文字很少。两个程序已用本书环境（小熊猫C++ 3.4 自带 GCC 11.5.0）编译通过，点餐程序已实际运行；源程序放入资源包。图中注明：手机点餐 App 的界面用其他技术制作，但“点菜—算钱—结账”的逻辑和这个程序一样。'],
    ['5', '①预处理指令段', '有没有视频或动画可以直观体现', '新做动画“预处理在做什么”（约 18 秒，MP4 和 GIF 各一份），放入资源包；书中新增静态分镜图1-9（预处理前—预处理后—编译后），本段末尾加“见图1-9”和动画说明。资源平台地址确定后可加二维码。'],
    ['6', '图1-8 执行流程（现图1-10）', '改成常用的这张图', '按所附图重画为竖向流程图：开始—编辑—编译—有错？—连接—执行—结果正确？—结束，标出 f.c、f.obj、f.exe 和“库函数和其他目标程序”。正文（1）～（3）的示例文件名同步改为 f.c、f.o 或 f.obj、f.exe。'],
    ['7', '项目小结', '用一段话总结项目，包括任务和知识点', '改为一段话，依次概括任务一、任务二各自做了什么、学了哪些知识点，以及 AI 升华；删去了不存在的“表1-14”。自查清单补“输出整数”和“四要素提问法”两项。'],
], [1.1, 2.6, 4.2, 8.7])

d.add_heading('三、编辑发现的问题与处理', level=1)
table(['位置', '问题', '处理'], [
    ['任务一“计算机语言”末段', '仍写“这个过程将在任务三中详细介绍”，合并后已没有任务三', '改为“任务二的‘C语言程序的执行流程’”'],
    ['全项目图号', '新增图1-6、图1-9', '原图1-6～1-11 依次顺延为图1-7～1-13，正文引用全部同步，已逐条核对'],
    ['全项目表号', '缺表1-3，表1-6 用了两次', '依次改为表1-3～表1-7'],
    ['任务一 任务实施', '“第7页的表1-2”“第6页的四条搜索话术”写死了页码，排版后会变', '改为“本任务后面的表1-2”“前面‘AI 助学 · 搜索话术’中的四条话术”'],
    ['“4. 小熊猫C++开发环境”“5. 规则与规范”', '用了下一级“1）”的标题样式，与 1～3 不同级', '改为与 1～3 同级'],
    ['例1-1 标题', '“欢迎使用校园AI迎新系统！。”双重标点', '删去句号'],
    ['例1-2 程序和表1-4', 'int a,b,sum; sum=a+b; 不符合本书自己提出的书写规范；几行注释过长，排版后折行', '改为 int a, b, sum;、sum = a + b; 等；缩短注释并对齐。例1-1 第 1 行注释同样缩短'],
    ['规则与规范 规范示例', '用了 %s，本项目没有讲', '改为直接输出“姓名：林小禾”'],
    ['表1-6 规则⑤', '错误示例只有 mian 一种，提示列却有两条（含 redefinition of \'main\'）', '示例补“或者写了两个 main 函数”'],
    ['执行流程（2）编译', '“目标程序默认存放在用户当前目录下”与小熊猫C++ 的实际情况不符', '改为：编译单个源文件时编译和连接一次完成，目标程序只是临时文件，文件夹中通常只看到 .c 和 .exe'],
    ['②main 主函数段', '括号里的话与前半句重复', '精简为“不论 main 主函数放在程序的什么位置（最前边、中间还是最后边）”'],
    ['小熊猫C++ 介绍', '“它由国内开发者……发展而来的开源软件”缺“是”', '补“是”'],
    ['计算机程序段', '“编辑好的指令代码”', '改为“编写好的指令”'],
    ['为什么学 C 语言 第 2 条', '“PLC 外围控制器”说法不准确（PLC 本身就是控制器；此处为 v4 原文遗留）', '改为“PLC（可编程控制器）”'],
    ['训练 1 slip.c', '行号从 6 跳到 11；“如下图所示”', '行号改为连续；改为“如图1-12所示”'],
    ['AI升华 profile8.c', '顺序与 profile.c 不同；“年龄：18 岁”与前文“年龄：18”不一致；“状态： 已报到”多一个空格', '改为 profile.c 原顺序加身高、籍贯、状态三行，与提问卡、AI 回答和运行结果一致'],
    ['实训巩固第 8 题', '#include <stdio.h>; 在本书环境中只有警告（已用 GCC 验证：extra tokens at end of #include directive），学生按“错误”找会困惑', '题干注明“其中 1 处编译器只给出[警告]”'],
    ['项目小结 自查清单', '“能用 printf 输出文字”漏了整数', '补“和整数”'],
    ['程序设计小节', '有不可见的零宽字符（从网页复制带入），影响检索和排版', '删除'],
    ['分页', '用连续空段落凑分页（任务二前 12 个，AI 使用小结前后 8 个），改稿后会错位', '删去空段落；任务二标题设为“段前分页”'],
], [3.3, 6.6, 6.7])

d.add_heading('四、请作者确认或提供', level=1)
p('图1-5 六张图片的来源如下。这些项目都在 GitHub 公开发布，许可证允许商业出版，条件是注明作者和许可证（书中图下已注明）。CC BY-SA 4.0 的两张经过裁剪，裁剪后的图仍按 CC BY-SA 4.0 授权。自动控制一图上有项目原带的英文标注，如嫌不便可换成实训室实拍。正式付印前，建议出版社按惯例再核对一次授权。', '1. 图1-5 图片来源与授权（请确认）。')
table(['位置', '图片内容', '来源（原图地址）', '许可证'], [
    ['自动控制', 'STM32 Nucleo 开发板＋SimpleFOCShield 驱动板＋无刷电机', 'SimpleFOC 项目文档 github.com/simplefoc/simplefoc.github.io（extras/Images/nucleo_foc_shield_connection.jpg）', 'MIT'],
    ['光伏发电', 'Libre Solar MPPT 2420 HC 光伏充电控制器电路板（设计：Martin Jäger）', 'github.com/LibreSolar/mppt-2420-hc（build/mppt-2420-hc_board.jpg）', 'CERN-OHL-W-2.0（仓库许可证）'],
    ['机器人', 'Thor 开源六轴机械臂（取宣传图左半部分）', 'github.com/AngelLM/Thor（doc/banner.png）', 'CC BY-SA 4.0'],
    ['无人机', 'Holybro X500 V2 四旋翼，装 Pixhawk 6C 飞控', 'PX4 自动驾驶仪文档 github.com/PX4/PX4-Autopilot（docs/assets/airframes/multicopter/x500_v2_holybro_pixhawk6c/kit_no_props.jpg）', 'CC BY 4.0'],
    ['装备制造', 'V1 Engineering LowRider 4 数控雕刻机（取中部）', 'github.com/V1EngineeringInc/V1EngineeringInc-Docs（docs/img/lr4/widelr4.jpg）', 'CC BY-SA 4.0'],
    ['系统软件', 'Linux 终端：查看内核、SQLite、Python 版本，用 gcc 编译运行 hello.c', '编者截图', '无第三方版权'],
], [1.8, 4.3, 7.7, 2.8])
p('维基共享资源、Pexels 等常用图库在本环境的网络中被拦截，所以这次只能从可访问的开源项目仓库中选图。没有采用的候选：PAROL6 机械臂（GPLv3，不适合图片出版）、某四旋翼照片（带商品包装盒）、LinuxCNC 文档配图（GFDL，出版条件较繁琐）。')
p('学校资源平台地址确定后，在图1-9 旁加动画二维码。', '2. 动画二维码。')
p('“项目八中的函数声明示例”和“安装过程请参照任务指导书中的任务一”引用的是后续项目和配套任务书，本轮无法核对，请确认编号。', '3. 引用核对。')
p('例1-1、例1-2、个人专属档案等运行结果截图是作者新绘的，本轮未改；Word 渲染与印刷排版字体不同，代码行尾注释的对齐以最终排版为准。', '4. 其他。')

d.add_heading('五、交付文件', level=1)
for line in ['项目一-搭建个人专属档案-修订稿v5（带批注回复）.docx：在作者稿上修改，保留 7 条批注并逐条回复',
             '项目一-搭建个人专属档案-修订稿v5（清洁版）.docx：同一内容，已删去全部批注',
             '项目一-搭建个人专属档案-修订稿v5-预览.pdf：清洁版的预览（LibreOffice 渲染，字体与 Word 略有不同）',
             '项目一-修订稿v5-配套资源.zip：预处理动画（MP4、GIF）、canteen.c 和 snake.c 源程序、新图高清原图、图1-5 六张图片原图及来源与许可说明']:
    d.add_paragraph(line, style='List Bullet')
d.save(OUT)
print('ok')
